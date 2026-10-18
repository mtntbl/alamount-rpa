# -*- coding: utf-8 -*-
"""
İşe giriş veri hazırlama (Mikro ERP - Personel Sicil Kartları, 5 ekran).
  Puantaj Excel'i (openpyxl) + personel PDF'i (OCR) → TC eşleştirme → checklist kuralları
  → girişe hazır, renk kodlu Excel + robotun okuyacağı JSON.
"""

import datetime as dt
import json
import logging
import os
import re
import shutil

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import belge_ocr as bo
import marka
from llm_katmani import LlmKatmani

log = logging.getLogger("alamount")

OTO, HR, YOK, BOS, LLM = "Otomatik", "HR ONAYI GEREKLİ", "Bilgi Yok", "Kasıtlı Boş", "LLM - KONTROL EDİN"

EKRANLAR = {
    1: "1. EKRAN - PERSONEL TANITIM KARTI",
    2: "2. EKRAN - PERSONEL İLETİŞİM BİLGİLERİ",
    3: "3. EKRAN - PERSONEL KİMLİK BİLGİLERİ",
    4: "4. EKRAN - PERSONEL NÜFUS BİLGİLERİ",
    5: "5. EKRAN - PERSONEL ÜCRET BİLGİLERİ",
}

# Puantaj sütun eşleştirmesi: anahtar -> başlıkta aranacak (sadeleştirilmiş) ifadeler
PUANTAJ_SUTUNLARI = {
    "tc": ["TC KIMLIK", "TCKN", "T.C", "KIMLIK NO"],
    "ad_soyad": ["ADI SOYADI", "AD SOYAD", "ADI-SOYADI"],
    "dogum_tarihi": ["DOGUM TARIHI"],
    "emekli": ["EMEKLI"],
    "birim": ["BIRIM", "TERSANE"],
    "giris_tarihi": ["ISE GIRIS", "GIRIS TARIHI"],
    "meslek": ["MESLEK"],
    "egitim": ["EGITIM", "TAHSIL"],
    "iban": ["IBN", "IBAN"],
    "telefon": ["TELEFON", "CEP"],
}

SICILLER = ["Ada", "Anadolu", "Gebze", "Tokat", "CMS-Sedef"]


# --------------------------------------------------------------------------- puantaj
def _hucre(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip() if v is not None else ""


def puantaj_oku(yol, sayfa="GİRİŞ"):
    """{tc: {anahtar: değer}} döner. Başlık satırını otomatik bulur."""
    wb = load_workbook(yol, data_only=True, read_only=True)
    ws = None
    for ad in wb.sheetnames:
        if bo.sadelestir(ad) == bo.sadelestir(sayfa):
            ws = wb[ad]
    if ws is None:
        log.warning("Puantaj'da '%s' sayfası yok, ilk sayfa kullanılıyor", sayfa)
        ws = wb[wb.sheetnames[0]]
    satirlar = list(ws.iter_rows(values_only=True))
    wb.close()

    baslik_i, eslesme = None, {}
    for i, row in enumerate(satirlar[:30]):
        basliklar = [bo.sadelestir(_hucre(c)) for c in row]
        es = {}
        for anahtar, ifadeler in PUANTAJ_SUTUNLARI.items():
            for j, b in enumerate(basliklar):
                if b and any(x in b for x in ifadeler) and j not in es.values():
                    es[anahtar] = j
                    break
        if "tc" in es and len(es) >= 3:
            baslik_i, eslesme = i, es
            break
    if baslik_i is None:
        raise ValueError("Puantaj'da 'TC KİMLİK NO' başlık satırı bulunamadı")

    kayitlar = {}
    for row in satirlar[baslik_i + 1:]:
        tc = re.sub(r"\D", "", _hucre(row[eslesme["tc"]]) if eslesme["tc"] < len(row) else "")
        if len(tc) != 11:
            continue
        kayitlar[tc] = {k: _hucre(row[j]) if j < len(row) else "" for k, j in eslesme.items()}
    log.info("Puantaj okundu: %d personel", len(kayitlar))
    return kayitlar


# --------------------------------------------------------------------------- kurallar
class Alan:
    __slots__ = ("ekran", "no", "ad", "deger", "durum", "not_", "kaynak", "anahtar")

    def __init__(self, ekran, no, ad, deger=None, durum=None, not_="", kaynak="", anahtar=None):
        self.ekran, self.no, self.ad = ekran, no, ad
        self.deger = deger if deger not in ("", None) else None
        self.durum = durum or (OTO if self.deger else YOK)
        self.not_, self.kaynak, self.anahtar = not_, kaynak, anahtar

    def sozluk(self):
        return {"ekran": self.ekran, "no": self.no, "alan": self.ad, "deger": self.deger or "",
                "durum": self.durum, "not": self.not_, "kaynak": self.kaynak}


def _ilk(*degerler):
    for kaynak, v in degerler:
        if v:
            return v, kaynak
    return None, ""


def _emekli_mi(v):
    s = bo.sadelestir(v)
    return "EMEKLI" in s and "DEGIL" not in s


def _iban(v):
    s = re.sub(r"\s", "", (v or "").upper())
    if not s:
        return None, ""
    if not s.startswith("TR"):
        s = "TR" + s
    if not re.fullmatch(r"TR\d{24}", s):
        return s, "IBAN biçimi beklenen TR + 24 hane değil, kontrol edin"
    sayi = "".join(str(int(c, 36)) for c in s[4:] + s[:4])
    return s, ("" if int(sayi) % 97 == 1 else "IBAN kontrol basamağı hatalı, kontrol edin")


def _telefon(v):
    d = re.sub(r"\D", "", v or "")
    if len(d) == 12 and d.startswith("90"):
        d = d[2:]
    if len(d) == 11 and d.startswith("0"):
        d = d[1:]
    return f"0{d[:3]} {d[3:6]} {d[6:8]} {d[8:]}" if len(d) == 10 else (v or None)


def alanlari_olustur(tc, p, belge, ayarlar):
    """Checklist'teki tüm alanları kural + puantaj (p) + belge bilgileriyle doldurur."""
    kim = belge["belgeler"].get("kimlik", {})
    nuf = belge["belgeler"].get("nufus_kayit", {})
    yer = belge["belgeler"].get("yerlesim", {})
    dip = belge["belgeler"].get("diploma", {})
    ser = belge["belgeler"].get("sertifika", {})

    birim = p.get("birim", "")
    sedef_listesi = [bo.sadelestir(x) for x in ayarlar.get("ISE_GIRIS", "sedef_birimleri").split(",") if x.strip()]
    sedef = any(s and s in bo.sadelestir(birim) for s in sedef_listesi)
    emekli = _emekli_mi(p.get("emekli", ""))

    sicil_oneri = "CMS-Sedef" if sedef else next(
        (s for s in SICILLER if bo.sadelestir(s) in bo.sadelestir(birim)), None)
    sicil_not = (f"Öneri: {sicil_oneri} (Birim={birim}). " if sicil_oneri else "") + "Şirket politikası, HR seçmeli."

    A = []
    E = lambda *a, **k: A.append(Alan(*a, **k))  # noqa: E731

    # 1. EKRAN
    E(1, 1, "Giriş Tarihi", p.get("giris_tarihi"), kaynak="Puantaj")
    E(1, 2, "Hangi Tersane?", birim, kaynak="Puantaj")
    E(1, 3, "Hangi Sicile Kayıt?", sicil_oneri, HR, sicil_not, "Kural (öneri)")
    E(1, 4, "Entegre Grubu", None, HR, "740 | 770 — belgelerden çıkarılamaz", "Şirket politikası")
    if emekli:
        E(1, 5, "Emekli mi / Normal mi / Huzur Hakkı", "Emekli", OTO, "", "Puantaj")
    else:
        E(1, 5, "Emekli mi / Normal mi / Huzur Hakkı", "Normal" if p.get("emekli") else None,
          OTO if p.get("emekli") else YOK, "EMEKLİ DEĞİL → Normal varsayıldı; Huzur Hakkı ihtimali varsa teyit edin",
          "Puantaj")

    # 2. EKRAN (güncel ikamet - Yerleşim Yeri Belgesi)
    adres_not = "" if yer.get("tam_adres") else "Yerleşim yeri belgesi bulunamadı/okunamadı"
    for no, ad, k in ((6, "Cadde", "cadde"), (7, "Mahalle", "mahalle"), (8, "Sokak", "sokak"),
                      (9, "Apartman / Bina No", "bina_no"), (10, "Daire No", "daire_no"),
                      (11, "Ülke", "ulke"), (12, "İl", "il"), (13, "İlçe", "ilce")):
        E(2, no, ad, yer.get(k) if yer.get("tam_adres") else None, not_=adres_not,
          kaynak="PDF - Yerleşim Yeri", anahtar=f"adres_{k}")
    E(2, 14, "Cep Telefonu", _telefon(p.get("telefon")), kaynak="Puantaj")

    # 3. EKRAN
    tahsil_not = ""
    if dip.get("tahsil") and p.get("egitim") and bo.sadelestir(dip["tahsil"]) not in bo.sadelestir(p["egitim"]):
        tahsil_not = f"Diploma '{dip['tahsil']}' diyor, Puantaj '{p['egitim']}' — kontrol edin"
    v, kay = _ilk(("Puantaj", p.get("egitim")), ("PDF - Diploma", dip.get("tahsil")))
    E(3, 15, "Tahsili", v, not_=tahsil_not, kaynak=kay)
    meslek_not, meslek_durum = "", None
    if ser.get("sertifika_meslek") and p.get("meslek") and \
            bo.sadelestir(ser["sertifika_meslek"]) != bo.sadelestir(p["meslek"]):
        meslek_not = f"Sertifika '{ser['sertifika_meslek']}' diyor, Puantaj '{p['meslek']}' — HR karar vermeli"
        meslek_durum = HR
    E(3, 16, "Mesleği", p.get("meslek"), meslek_durum, meslek_not, "Puantaj")
    E(3, 17, "Görevi", None, HR, f"Sayısal görev kodu (örn. 7213.32) gerekli; Puantaj: '{p.get('meslek', '')}'",
      "HR")
    E(3, 18, "SGK (Hangi Sicile Kayıt)", sicil_oneri, HR, "Alan 3 ile aynı olmalı", "Kural (öneri)")
    ana = "02" if emekli else "01"
    E(3, 19, "Tabi olduğu Ana Kanun", ana, OTO, "Emekli → 02, değil → 01", "Kural")
    if sedef:
        k2, k2_not = "0000", f"SEDEF İSTİSNASI uygulandı (Birim={birim})"
    else:
        k2, k2_not = ("0000", "Emekli → 0000") if emekli else ("5510", "Emekli değil → 5510")
    E(3, 20, "Tabi olduğu kanun 2", k2, OTO, k2_not, "Kural")
    E(3, 21, "Tabi olduğu diğer belge türü", ana, OTO, "Alan 19 ile aynı mantık", "Kural")
    E(3, 22, "Tabi olduğu diğer kanun", k2, OTO, k2_not, "Kural")

    # 4. EKRAN
    yabanci = tc.startswith("99")
    uyruk, kay = _ilk(("PDF - Kimlik", kim.get("uyruk")), ("TC numarası", None if yabanci else "T.C."))
    E(4, 23, "Uyruğu", uyruk, not_="Yabancı kimlik no (99…) — HR kontrol" if yabanci else "", kaynak=kay,
      anahtar="uyruk")
    E(4, 24, "Yabancı uyruklu ise ülke adı", None, BOS if not yabanci else YOK, "", "", anahtar=None)
    v, kay = _ilk(("PDF - Kimlik", kim.get("cinsiyet")), ("PDF - Nüfus Kayıt", nuf.get("cinsiyet")))
    E(4, 25, "Cinsiyeti", v, kaynak=kay, anahtar="cinsiyet")
    E(4, 26, "Medeni Hal", nuf.get("medeni_hal"), kaynak="PDF - Nüfus Kayıt", anahtar="medeni_hal")
    E(4, 27, "Dini", nuf.get("din"), kaynak="PDF - Nüfus Kayıt", anahtar="din")
    v, kay = _ilk(("PDF - Kimlik", kim.get("dogum_tarihi")), ("PDF - Nüfus Kayıt", nuf.get("dogum_tarihi")),
                  ("Puantaj", p.get("dogum_tarihi")))
    dt_not = ""
    if v and p.get("dogum_tarihi") and v != p["dogum_tarihi"]:
        dt_not = f"Puantaj'daki doğum tarihi farklı: {p['dogum_tarihi']}"
    E(4, 28, "Doğum tarihi", v, HR if dt_not else None, dt_not, kay, anahtar="dogum_tarihi")
    E(4, 29, "Doğum yeri", nuf.get("dogum_yeri"), kaynak="PDF - Nüfus Kayıt", anahtar="dogum_yeri")
    E(4, 30, "Kan Grubu", None, YOK, "Belgelerde genelde bulunmaz", "", anahtar=None)
    E(4, 31, "Seri no", kim.get("seri_no"), kaynak="PDF - Kimlik", anahtar="seri_no")
    E(4, 32, "İl (Nüfus)", nuf.get("il"), kaynak="PDF - Nüfus Kayıt", anahtar="nufus_il")
    E(4, 33, "İlçe", nuf.get("ilce"), kaynak="PDF - Nüfus Kayıt", anahtar="nufus_ilce")
    E(4, 34, "Mahalle", nuf.get("mahalle"), kaynak="PDF - Nüfus Kayıt", anahtar="nufus_mahalle")
    E(4, 35, "Köy", nuf.get("koy"), kaynak="PDF - Nüfus Kayıt", anahtar=None)
    E(4, 36, "Cilt No", nuf.get("cilt_no"), kaynak="PDF - Nüfus Kayıt", anahtar="cilt_no")
    E(4, 37, "Aile Sıra No", nuf.get("aile_sira_no"), not_="Hane No ile aynı olabilir, teyit edin",
      kaynak="PDF - Nüfus Kayıt", anahtar="aile_sira_no")
    E(4, 38, "Sıra no", nuf.get("sira_no"), kaynak="PDF - Nüfus Kayıt", anahtar="sira_no")
    eski = "Eski nüfus cüzdanına özgü; yeni çipli kartta basılı değil"
    E(4, 39, "Veriliş Nedeni", None, YOK, eski, "")
    E(4, 40, "Verildiği nüfus idaresi", kim.get("veren_makam"), kaynak="PDF - Kimlik (arka)", anahtar="veren_makam")
    E(4, 41, "Veriliş tarihi", None, YOK, eski, "")
    E(4, 42, "Cüzdan Kayıt No", None, YOK, eski + " (yeni kartta PEN no olabilir)", "")
    pdf_tc = belge.get("tc_listesi", [])
    if not pdf_tc:
        E(4, 43, "T.C. kimlik no", tc, HR, "PDF'te TC bulunamadı, çapraz doğrulama yapılamadı", "Puantaj")
    elif tc in pdf_tc and len(pdf_tc) == 1:
        E(4, 43, "T.C. kimlik no", tc, OTO, "Puantaj ve PDF eşleşti", "Puantaj + PDF")
    elif tc in pdf_tc:
        E(4, 43, "T.C. kimlik no", tc, HR, f"PDF'te başka TC'ler de var: {', '.join(map(bo.tc_maskele, pdf_tc))}",
          "Puantaj + PDF")
    else:
        E(4, 43, "T.C. kimlik no", tc, HR, "PDF'teki TC Puantaj ile EŞLEŞMİYOR — yanlış belge olabilir!",
          "Puantaj + PDF")
    E(4, 44, "İlk soyad", nuf.get("ilk_soyad"), BOS if not nuf.get("ilk_soyad") else None,
      "Varsa önceki soyad", "PDF - Nüfus Kayıt", anahtar=None)
    E(4, 45, "Orijinal dilinde adı soyadı", None, BOS if not yabanci else YOK, "Yabancılar için", "")

    # 5. EKRAN
    E(5, 46, "Ücret tipi", "Aylık", OTO, "Her zaman Aylık", "Kural")
    E(5, 47, "Ücret", None, BOS, "KASITLI BOŞ — yetkili kişi Mikro'ya kendisi girer", "Kural")
    E(5, 48, "Hesap tipi", "Diğer", OTO, "Her zaman Diğer", "Kural")
    E(5, 49, "SGK Şubesi", None, HR, "Alan 3/18 ile senkron olmalı" + (f" (öneri: {sicil_oneri})" if sicil_oneri else ""),
      "Şirket politikası")
    E(5, 50, "Varsa özel SGK değer kodu", None, BOS, "Boş bırakılır", "Kural")
    E(5, 51, "Çalışma kodu", "0", OTO, "Genelde 0", "Kural")
    E(5, 52, "Meslek kodu", None, HR, f"Puantaj'da yalnızca metin var: '{p.get('meslek', '')}'", "HR")
    E(5, 54, "Banka ödeme", None, HR, "Yapılsın / Yapılmasın — şirket politikası", "Şirket politikası")
    iban, iban_not = _iban(p.get("iban"))
    E(5, 55, "Hesap / IBAN no", iban, HR if iban_not else None, iban_not, "Puantaj")
    return A


LLM_ALAN_ACIKLAMALARI = {
    "adres_mahalle": "Yerleşim yeri adresindeki mahalle adı (MAH. olmadan)",
    "adres_cadde": "Yerleşim yeri adresindeki cadde adı",
    "adres_sokak": "Yerleşim yeri adresindeki sokak adı",
    "adres_bina_no": "Dış kapı / bina no",
    "adres_daire_no": "İç kapı / daire no",
    "adres_il": "Yerleşim yeri il", "adres_ilce": "Yerleşim yeri ilçe", "adres_ulke": "Ülke",
    "uyruk": "Uyruğu (T.C. ya da ülke)",
    "cinsiyet": "Cinsiyet: Erkek ya da Kadın",
    "medeni_hal": "Medeni hal: Bekar/Evli/Dul/Boşanmış",
    "din": "Dini (nüfus kayıt örneği)",
    "dogum_tarihi": "Doğum tarihi gg.aa.yyyy",
    "dogum_yeri": "Doğum yeri",
    "seri_no": "Kimlik kartı seri no (örn. A12B34567)",
    "nufus_il": "Nüfusa kayıtlı il", "nufus_ilce": "Nüfusa kayıtlı ilçe",
    "nufus_mahalle": "Nüfusa kayıtlı mahalle/köy",
    "cilt_no": "Cilt no", "aile_sira_no": "Aile sıra no / hane no", "sira_no": "Birey sıra no (BSN)",
    "veren_makam": "Kimliği veren makam / nüfus idaresi",
}


def llm_ile_tamamla(alanlar, belge, ayarlar):
    llm = LlmKatmani(ayarlar)
    if not llm.aktif:
        return 0
    sadece_eksik = ayarlar.getbool("LLM", "sadece_eksikler")
    hedef = {a.anahtar: a for a in alanlar
             if a.anahtar in LLM_ALAN_ACIKLAMALARI and (a.durum == YOK or not sadece_eksik)}
    if not hedef:
        return 0
    metinler = [f"[Sayfa {s['no']} - {bo.BELGE_ADLARI.get(s['tur'], s['tur'])}]\n{s['metin']}"
                for s in belge["sayfalar"] if s["tur"] in ("kimlik", "nufus_kayit", "yerlesim", "diger")]
    sonuc = llm.alanlari_cikar({k: LLM_ALAN_ACIKLAMALARI[k] for k in hedef}, metinler,
                               belge.get("kimlik_gorselleri", []))
    n = 0
    for k, v in sonuc.items():
        a = hedef[k]
        if v and v != a.deger:
            if a.deger:
                a.not_ = (a.not_ + " | " if a.not_ else "") + f"LLM farklı okudu: {v}"
                a.durum = HR
            else:
                a.deger, a.durum, a.kaynak = v, LLM, "LLM (Claude)"
                n += 1
    return n


# --------------------------------------------------------------------------- Excel çıktısı
RENK = {
    "baslik": PatternFill("solid", fgColor="022351"),
    "ekran": PatternFill("solid", fgColor="DDF4FD"),
    HR: PatternFill("solid", fgColor="FFF2CC"),
    YOK: PatternFill("solid", fgColor="D9D9D9"),
    BOS: PatternFill("solid", fgColor="EDEDED"),
    LLM: PatternFill("solid", fgColor="FCE4D6"),
}


def excel_yaz(yol, kisi, alanlar, belge):
    wb = Workbook()
    ws = wb.active
    ws.title = "Mikro Giriş"
    ince = Side(style="thin", color="BFBFBF")
    kenar = Border(left=ince, right=ince, top=ince, bottom=ince)

    wb.properties.creator = marka.URUN
    ws.append([f"İşe Giriş - {kisi['ad_soyad']} ({kisi['tc']})"])
    ws["A1"].font = Font(bold=True, size=14, color="022351")
    ws.append([f"Hazırlanma: {dt.datetime.now():%d.%m.%Y %H:%M}  |  Birim: {kisi.get('birim', '')}  |  {marka.URUN}"])
    ws["A2"].font = Font(color="5F6B7A", size=9)
    ws.row_dimensions[1].height = 30
    try:  # sağ üst köşeye logo
        from openpyxl.drawing.image import Image as XlResim

        logo = XlResim(marka.kaynak("logo.png"))
        logo.width, logo.height = 150, int(150 * logo.height / logo.width)
        ws.add_image(logo, "E1")
    except Exception as hata:
        log.debug("Excel'e logo eklenemedi: %s", hata)
    ws.append([])
    ws.append(["No", "Alan Adı", "Değer", "Durum", "Not", "Kaynak"])
    for c in ws[4]:
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), RENK["baslik"], kenar

    ekran = None
    for a in alanlar:
        if a.ekran != ekran:
            ekran = a.ekran
            ws.append([EKRANLAR[ekran]])
            r = ws.max_row
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            ws.cell(r, 1).font = Font(bold=True, color="022351")
            for col in range(1, 7):
                ws.cell(r, col).fill = RENK["ekran"]
        ws.append([a.no, a.ad, a.deger or "", a.durum, a.not_, a.kaynak])
        r = ws.max_row
        for col in range(1, 7):
            c = ws.cell(r, col)
            c.border = kenar
            c.alignment = Alignment(vertical="top", wrap_text=col in (2, 3, 5))
            if a.durum in RENK:
                c.fill = RENK[a.durum]
    for col, w in zip("ABCDEF", (6, 34, 30, 20, 60, 22)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"

    oz = wb.create_sheet("Özet")
    say = {}
    for a in alanlar:
        say[a.durum] = say.get(a.durum, 0) + 1
    oz.append(["Durum", "Alan sayısı"])
    for k, v in say.items():
        oz.append([k, v])
    oz.append([])
    oz.append(["Sayfa", "Belge türü", "Okuma", "Kırpılan kart"])
    for s in belge["sayfalar"]:
        oz.append([s["no"], bo.BELGE_ADLARI.get(s["tur"], s["tur"]), "OCR" if s["kaynak"] == "ocr" else "Metin katmanı",
                   s.get("kart_sayisi", 0)])
    for row in (oz[1], oz[len(say) + 3]):
        for c in row:
            c.font = Font(bold=True)
    oz.column_dimensions["A"].width = 22
    oz.column_dimensions["B"].width = 26
    oz.column_dimensions["C"].width = 16
    wb.save(yol)


def guvenli_ad(s):
    return re.sub(r'[<>:"/\\|?*]+', "_", s).strip() or "PERSONEL"


# --------------------------------------------------------------------------- ana akış
def hazirla(tc, pdf_yolu, puantaj_yolu, ayarlar, cikti_kok=None, puantaj=None):
    """Tek personel için girişe hazır Excel + JSON üretir. Sonuç sözlüğü döner."""
    tc = re.sub(r"\D", "", str(tc))
    gizle = bo.tc_maskele if ayarlar.getbool("ISE_GIRIS", "tc_maskele") else (lambda x: x)
    if not bo.tc_gecerli_mi(tc):
        raise ValueError(f"Geçersiz TC Kimlik No: {gizle(tc)}")
    puantaj = puantaj if puantaj is not None else puantaj_oku(puantaj_yolu, ayarlar.get("ISE_GIRIS", "puantaj_sayfa"))
    if tc not in puantaj:
        raise LookupError(f"TC {gizle(tc)} Puantaj'da bulunamadı")
    p = puantaj[tc]
    ad = p.get("ad_soyad") or "PERSONEL"
    log.info("İşe giriş hazırlanıyor: %s", gizle(tc))

    cikti_kok = cikti_kok or ayarlar.yol("ISE_GIRIS", "cikti", olustur=True)
    klasor = os.path.join(cikti_kok, guvenli_ad(f"{ad} ({tc})") + f"_{dt.datetime.now():%Y%m%d_%H%M%S}")
    os.makedirs(klasor, exist_ok=True)

    belge = bo.pdf_coz(pdf_yolu, ayarlar, os.path.join(klasor, "gorseller"))
    alanlar = alanlari_olustur(tc, p, belge, ayarlar)
    n_llm = llm_ile_tamamla(alanlar, belge, ayarlar)

    kisi = {"tc": tc, "ad_soyad": ad, "birim": p.get("birim", "")}
    xlsx = os.path.join(klasor, guvenli_ad(f"İşe Girişi Hazır - {ad} ({tc})") + ".xlsx")
    excel_yaz(xlsx, kisi, alanlar, belge)
    js = os.path.join(klasor, "hazir_veri.json")
    with open(js, "w", encoding="utf-8") as f:
        json.dump({"kisi": kisi, "alanlar": [a.sozluk() for a in alanlar]}, f, ensure_ascii=False, indent=2)

    ozet = {}
    for a in alanlar:
        ozet[a.durum] = ozet.get(a.durum, 0) + 1
    uyarilar = [f"{a.no}. {a.ad}: {a.not_}" for a in alanlar if a.durum == HR and a.not_ and "öneri" not in a.not_.lower()
                and "politika" not in a.not_.lower()]
    log.info("Hazır: %s | %s%s", os.path.basename(xlsx), ", ".join(f"{k}={v}" for k, v in ozet.items()),
             f" | LLM {n_llm} alan" if n_llm else "")
    return {"excel": xlsx, "json": js, "klasor": klasor, "ozet": ozet, "uyarilar": uyarilar, "ad_soyad": ad, "tc": tc}


def klasoru_isle(ayarlar):
    """Giriş klasöründeki <TC>*.pdf dosyalarını toplu işler."""
    giris = ayarlar.yol("ISE_GIRIS", "giris", olustur=True)
    puantaj_yolu = ayarlar.yol("ISE_GIRIS", "puantaj")
    pdfler = [f for f in sorted(os.listdir(giris)) if f.lower().endswith(".pdf")]
    if not pdfler:
        log.info("İşe giriş: işlenecek PDF yok (%s)", giris)
        return []
    if not os.path.isfile(puantaj_yolu):
        raise FileNotFoundError(f"Puantaj dosyası bulunamadı: {puantaj_yolu}")
    puantaj = puantaj_oku(puantaj_yolu, ayarlar.get("ISE_GIRIS", "puantaj_sayfa"))
    sonuclar = []
    for f in pdfler:
        yol = os.path.join(giris, f)
        m = re.search(r"(?<!\d)(\d{11})(?!\d)", f)
        try:
            tc = m.group(1) if m else None
            if not tc:  # dosya adında yoksa PDF içinden bul
                belge_tc = bo.tc_bul(bo.pdf_metne_cevir(yol, ayarlar))
                tc = next((t for t in belge_tc if t in puantaj), None)
                if not tc:
                    raise LookupError("Dosya adında ve PDF içinde Puantaj'la eşleşen TC bulunamadı")
            sonuc = hazirla(tc, yol, puantaj_yolu, ayarlar, puantaj=puantaj)
            sonuclar.append(sonuc)
            if ayarlar.getbool("ISE_GIRIS", "arsive_tasi"):
                hedef = os.path.join(ayarlar.folder("arsiv"), "ise_giris")
                os.makedirs(hedef, exist_ok=True)
                shutil.move(yol, os.path.join(hedef, f))
        except Exception as e:
            log.error("İşe giriş HATA (%s): %s", f, e)
            hedef = os.path.join(ayarlar.folder("hata"), "ise_giris")
            os.makedirs(hedef, exist_ok=True)
            shutil.move(yol, os.path.join(hedef, f))
            with open(os.path.join(hedef, f + ".hata.txt"), "w", encoding="utf-8") as h:
                h.write(str(e))
            sonuclar.append({"hata": str(e), "dosya": f})
    return sonuclar

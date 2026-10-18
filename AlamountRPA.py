# -*- coding: utf-8 -*-
"""
Robot Framework anahtar kelime kütüphanesi.
.robot görevlerinde:   Library    AlamountRPA
Bot katmanlarını (PCA, OCR, işe giriş) Robot görevlerinden çağrılabilir hale getirir.
"""

import glob
import json
import logging
import os

from robot.api import logger
from robot.api.deco import keyword, library
from openpyxl import load_workbook

import ayarlar as ay
import belge_ocr as bo
import ise_giris
import pca_analiz


class _RobotLogHandler(logging.Handler):
    def emit(self, record):
        msg = self.format(record)
        if record.levelno >= logging.WARNING:
            logger.warn(msg)
        else:
            logger.info(msg)


@library(scope="GLOBAL", auto_keywords=False)
class AlamountRPA:
    def __init__(self, config=None):
        self.config = config or os.environ.get("ALAMOUNT_RPA_CONFIG") or ay.varsayilan_config_yolu()
        self.ayarlar = ay.ayarlari_yukle(self.config)
        log = logging.getLogger("alamount")
        if not any(isinstance(h, _RobotLogHandler) for h in log.handlers):
            h = _RobotLogHandler()
            h.setFormatter(logging.Formatter("%(message)s"))
            log.addHandler(h)
            log.setLevel(logging.INFO)

    # ------------------------------------------------------------------ ayarlar
    @keyword("Ayarı Oku")
    def ayari_oku(self, bolum, anahtar):
        """config.ini'den bir değer okur. Örnek: ``${mod}=  Ayarı Oku  GENEL  mod``"""
        return self.ayarlar.get(bolum, anahtar)

    @keyword("Ayar Evet Mi")
    def ayar_evet_mi(self, bolum, anahtar):
        return self.ayarlar.getbool(bolum, anahtar)

    @keyword("Klasör Yolu")
    def klasor_yolu(self, bolum, anahtar, olustur=True):
        """config.ini'deki göreli yolu exe klasörüne göre mutlak yola çevirir."""
        return self.ayarlar.yol(bolum, anahtar, olustur=bool(olustur))

    @keyword("Kuru Çalışma Mı")
    def kuru_calisma_mi(self, deger=None):
        """Değişken verilmişse onu, verilmemişse [ROBOT] kuru_calisma ayarını döner."""
        if deger not in (None, "", "${NONE}", "None"):
            return str(deger).strip().lower() in ("1", "true", "evet", "yes", "e")
        return self.ayarlar.getbool("ROBOT", "kuru_calisma")

    # ------------------------------------------------------------------ PCA
    @keyword("PCA Klasörünü İşle")
    def pca_klasorunu_isle(self):
        """input klasöründeki veri dosyalarına PCA uygular. Başarılı ve hatalı sayısını döner."""
        ok, hata = pca_analiz.run_once(self.ayarlar)
        return [ok, hata]

    @keyword("PCA Dosyası İşle")
    def pca_dosyasi_isle(self, dosya):
        return pca_analiz.process_file(os.path.abspath(dosya), self.ayarlar)

    # ------------------------------------------------------------------ belge / OCR
    @keyword("PDF Metne Çevir")
    def pdf_metne_cevir(self, pdf, txt=None):
        """PDF'i metne çevirir; taranmış sayfalara otomatik OCR uygular."""
        return bo.pdf_metne_cevir(pdf, self.ayarlar, txt)

    @keyword("PDF Sayfalarını Görsele Çevir")
    def pdf_sayfalarini_gorsele_cevir(self, pdf, klasor, dpi=None):
        os.makedirs(klasor, exist_ok=True)
        yollar = []
        for no, _, img in bo.pdf_sayfalari(pdf, int(dpi or self.ayarlar.getint("OCR", "dpi"))):
            p = os.path.join(klasor, f"sayfa{no:02d}.png")
            img.save(p)
            yollar.append(p)
        return yollar

    @keyword("Kimlik Kartlarını Kırp")
    def kimlik_kartlarini_kirp(self, gorsel, klasor):
        """Görseldeki kimlik kartı(larını) bulur, düzeltir ve PNG olarak kaydeder."""
        from PIL import Image

        os.makedirs(klasor, exist_ok=True)
        ad = os.path.splitext(os.path.basename(gorsel))[0]
        yollar = []
        for i, kart in enumerate(bo.kimlik_karti_kirp(Image.open(gorsel)), 1):
            p = os.path.join(klasor, f"{ad}_kimlik{i}.png")
            kart.save(p)
            yollar.append(p)
        return yollar

    @keyword("Görselde OCR Yap")
    def gorselde_ocr_yap(self, gorsel, psm=None):
        from PIL import Image

        return bo.OcrMotoru(self.ayarlar).oku(Image.open(gorsel), int(psm) if psm else None)

    @keyword("Metindeki TC Numaraları")
    def metindeki_tc_numaralari(self, metin):
        return bo.tc_bul(metin)

    @keyword("TC Geçerli Mi")
    def tc_gecerli_mi(self, tc):
        return bo.tc_gecerli_mi(str(tc))

    # ------------------------------------------------------------------ işe giriş
    @keyword("İşe Giriş Verisi Hazırla")
    def ise_giris_verisi_hazirla(self, tc, pdf, puantaj=None):
        """Tek personel için girişe hazır Excel + JSON üretir; sonuç sözlüğü döner."""
        puantaj = puantaj or self.ayarlar.yol("ISE_GIRIS", "puantaj")
        return ise_giris.hazirla(tc, pdf, puantaj, self.ayarlar)

    @keyword("İşe Giriş Klasörünü İşle")
    def ise_giris_klasorunu_isle(self):
        return ise_giris.klasoru_isle(self.ayarlar)

    @keyword("Son Hazır Excel")
    def son_hazir_excel(self):
        kok = self.ayarlar.yol("ISE_GIRIS", "cikti", olustur=True)
        dosyalar = glob.glob(os.path.join(kok, "*", "İşe Girişi Hazır*.xlsx"))
        if not dosyalar:
            raise FileNotFoundError(f"Hazır Excel bulunamadı: {kok}")
        return max(dosyalar, key=os.path.getmtime)

    @keyword("Excelden Giriş Verisi Oku")
    def excelden_giris_verisi_oku(self, xlsx):
        """
        HR'ın kontrol ettiği 'Mikro Giriş' sayfasını okur ve Mikro'ya girilecek alanları döner.
        Kurallar:
          * Değeri boş satırlar ve 'Kasıtlı Boş' satırlar girilmez.
          * 'HR ONAYI GEREKLİ' / 'LLM - KONTROL EDİN' satırları, HR Durum hücresini 'ONAYLANDI'
            yapmadıkça girilmez (uyarı yazılır).
        Döner: [{no, ekran, alan, deger}]
        """
        wb = load_workbook(xlsx, data_only=True)
        ws = wb["Mikro Giriş"]
        ekran, sonuc, atlanan = 0, [], []
        for row in ws.iter_rows(min_row=5, values_only=True):
            no, alan, deger, durum = (list(row) + [None] * 4)[:4]
            if isinstance(no, str) and "EKRAN" in no:
                ekran = int(no.split(".")[0])
                continue
            if not isinstance(no, (int, float)):
                continue
            deger = "" if deger is None else str(deger).strip()
            durum = (durum or "").strip()
            if not deger or durum == ise_giris.BOS:
                continue
            if durum in (ise_giris.HR, ise_giris.LLM):
                atlanan.append(f"{int(no)}. {alan}")
                continue
            sonuc.append({"no": int(no), "ekran": ekran, "alan": alan, "deger": deger})
        if atlanan:
            logger.warn("Onaylanmamış alanlar atlandı (Durum'u 'ONAYLANDI' yapın): " + ", ".join(atlanan))
        logger.info(f"{len(sonuc)} alan girilecek ({os.path.basename(xlsx)})", also_console=True)
        return sonuc

    # ------------------------------------------------------------------ Kontrol Merkezi
    def _merkez(self):
        from merkez_istemci import MerkezIstemci

        if not hasattr(self, "_istemci"):
            self._istemci = MerkezIstemci.ayarlardan(self.ayarlar)
        return self._istemci

    @keyword("Varlık Al")
    def varlik_al(self, ad, alan="deger"):
        """Merkezin şifreli kasasından varlık okur. ``alan``: deger | kullanici_adi.
        Örnek: ``${parola}=  Varlık Al  MIKRO_GIRIS`` / ``${kul}=  Varlık Al  MIKRO_GIRIS  kullanici_adi``"""
        v = self._merkez().varlik(ad)
        logger.info(f"Varlık okundu: {ad}")
        return v[alan]

    @keyword("Kuyruğa Ekle")
    def kuyruga_ekle(self, kuyruk, referans, veri=None, oncelik=5, max_deneme=3):
        r = self._merkez().kuyruga_ekle(kuyruk, referans, veri or {}, oncelik=int(oncelik), max_deneme=int(max_deneme))
        return bool(r and r.get("eklendi"))

    @keyword("Kuyruktan Öğe Al")
    def kuyruktan_oge_al(self, kuyruk):
        """Sıradaki öğeyi alır; kuyruk boşsa ``None`` döner. Öğe: id, referans, veri, deneme, son_deneme."""
        return self._merkez().kuyruktan_al(kuyruk)

    @keyword("Öğeyi Bitir")
    def ogeyi_bitir(self, oge, durum="basarili", hata="", sonuc=None):
        """durum: basarili | is_hatasi | sistem_hatasi"""
        oge_id = oge["id"] if isinstance(oge, dict) else int(oge)
        return self._merkez().oge_bitir(oge_id, durum, sonuc or {}, str(hata))

    @keyword("Hazır Veriyi Yükle")
    def hazir_veriyi_yukle(self, json_yolu):
        with open(json_yolu, encoding="utf-8") as f:
            return json.load(f)

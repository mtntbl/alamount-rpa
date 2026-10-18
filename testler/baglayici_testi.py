# -*- coding: utf-8 -*-
"""
Türkiye bağlayıcıları testi: VKN/TCKN/IBAN, e-Fatura (UBL-TR XML/ZIP), MT940, mutabakat, Excel, TCMB.
Kullanım: python testler/baglayici_testi.py <proje_klasoru> [--canli]   (--canli: TCMB'den gerçek kur çeker)
"""
import os
import sys
import tempfile
import zipfile

KOK = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else ".")
sys.path.insert(0, KOK)
import AlamountBaglayicilar as b  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

SONUC = {}
EF = os.path.join(KOK, "ornek_veri", "efatura")
MT = os.path.join(KOK, "ornek_veri", "banka", "ornek_ekstre.sta")
G = tempfile.mkdtemp()

# ---- doğrulama
SONUC["VKN geçerli/geçersiz"] = b.vkn_gecerli_mi("1234567890") and not b.vkn_gecerli_mi("1234567891") \
    and not b.vkn_gecerli_mi("12345")
SONUC["TCKN geçerli/geçersiz"] = b.tckn_gecerli_mi("10000000146") and not b.tckn_gecerli_mi("10000000145") \
    and not b.tckn_gecerli_mi("01234567890")
SONUC["IBAN geçerli/geçersiz"] = b.iban_gecerli_mi("TR33 0006 1005 1978 6457 8413 26") and \
    not b.iban_gecerli_mi("TR330006100519786457841327") and not b.iban_gecerli_mi("TR3300061005")
SONUC["Türkçe sayı biçimleri"] = [b._ondalik(x) for x in ("1.234,56", "1234.56", "1,234.56", "1234,5")] == \
    [b.Decimal("1234.56")] * 3 + [b.Decimal("1234.5")]

# ---- e-Fatura
f, hatalar = b.efatura_klasoru_oku(EF)
SONUC["3 fatura okundu (2 XML + 1 ZIP)"] = len(f) == 3 and not hatalar
f1 = next(x for x in f if x["fatura_no"] == "ORN2026000000123")
SONUC["fatura alanları"] = (f1["senaryo"], f1["tip"], f1["tarih"], f1["satici"]["vkn_tckn"],
                            f1["satici"]["vergi_dairesi"], f1["alici"]["ad"]) == \
    ("TICARIFATURA", "SATIS", "2026-10-02", "4820156735", "Kadıköy", "Alamount Teknoloji A.Ş.")
SONUC["fatura tutarları (KDV dahil)"] = (f1["mal_hizmet_toplam"], f1["kdv_toplam"], f1["odenecek_tutar"]) == \
    (20200.0, 4040.0, 24240.0)
SONUC["fatura satırları"] = len(f1["satirlar"]) == 2 and f1["satirlar"][1]["urun"] == "Toner Kartuşu" and \
    f1["satirlar"][1]["kdv_orani"] == 20.0
f3 = next(x for x in f if x["fatura_no"] == "ORN2026000000125")
SONUC["döviz faturası (USD + kur)"] = f3["para_birimi"] == "USD" and f3["kur"] == 41.921 and \
    f3["kaynak"].endswith(".zip:ORN2026000000125.xml")
SONUC["tutarlı faturada uyarı yok"] = all(not x["uyarilar"] for x in f)
bozuk = open(os.path.join(EF, "ORN2026000000124.xml"), encoding="utf-8").read() \
    .replace("<cbc:ID schemeID=\"VKN\">7301942583", "<cbc:ID schemeID=\"VKN\">7301942584") \
    .replace(">12500.00</cbc:LineExtensionAmount>", ">12000.00</cbc:LineExtensionAmount>", 1)
x = b.efatura_xml_coz(bozuk.encode("utf-8"), "bozuk.xml")
SONUC["tutarsız satır ve geçersiz VKN uyarısı"] = any("Satır toplamı" in u for u in x["uyarilar"]) and \
    any("VKN" in u for u in x["uyarilar"])
k2 = os.path.join(G, "karisik")
os.makedirs(k2)
open(os.path.join(k2, "bozuk.xml"), "w").write("<Invoice><eksik")
open(os.path.join(k2, "irsaliye.xml"), "w").write("<DespatchAdvice/>")
with zipfile.ZipFile(os.path.join(k2, "toplu.zip"), "w") as z:
    for ad in os.listdir(EF):
        if ad.endswith(".xml"):
            z.write(os.path.join(EF, ad), ad)
f_k, h_k = b.efatura_klasoru_oku(k2)
SONUC["bozuk dosya işi durdurmaz (toplu ZIP okunur)"] = len(f_k) == 2 and len(h_k) == 2

# ---- MT940
e = b.mt940_oku(MT)[0]
SONUC["MT940: hesap, bakiye, 3 hareket"] = e["hesap"] == "TR330006100519786457841326" and \
    e["acilis"]["tutar"] == 250000.0 and len(e["hareketler"]) == 3 and e["tutarli"] is True
h = e["hareketler"]
SONUC["MT940: yön, tutar, çok satırlı açıklama"] = h[0]["yon"] == "borç" and h[0]["tutar"] == -24240.0 and \
    h[1]["yon"] == "alacak" and "EKIM BAKIM BEDELI" in h[2]["aciklama"]
metin = open(MT, encoding="utf-8").read().replace("MUSTERI A.S.", "MÜŞTERİ A.Ş.")
p = os.path.join(G, "cp1254.sta")
open(p, "w", encoding="cp1254").write(metin)
SONUC["MT940: Windows-1254 kodlama"] = "MÜŞTERİ" in b.mt940_oku(p)[0]["hareketler"][1]["aciklama"]
SONUC["MT940: bakiye tutarsızlığı yakalanır"] = b.mt940_coz(metin.replace(":62F:C261010", ":62F:C261011")
                                                            .replace(",00\n-", "1,00\n-"))[0]["tutarli"] is False

# ---- mutabakat
s, u = b.faturalari_eslestir(f, e["hareketler"])
d = {r["fatura_no"]: (r["durum"], r["eslesme"]) for r in s}
SONUC["mutabakat: fatura no ile eşleşme"] = d["ORN2026000000123"] == ("ödendi", "fatura no")
SONUC["mutabakat: tutarla eşleşme"] = d["ORN2026000000124"] == ("ödendi", "tutar")
SONUC["mutabakat: ödenmemiş fatura"] = d["ORN2026000000125"][0] == "ödenmedi" and not u

# ---- Robot kütüphanesi + Excel
kut = b.AlamountBaglayicilar()
xl = kut.e_faturalari_excele_aktar(EF, os.path.join(G, "faturalar.xlsx"))
wb = load_workbook(xl)
SONUC["Excel: Faturalar + Satırlar sayfaları"] = wb.sheetnames == ["Faturalar", "Satırlar"] and \
    wb["Faturalar"].max_row == 4 and wb["Satırlar"].max_row == 6 and wb["Faturalar"]["P2"].value == 24240.0
mx = os.path.join(G, "mutabakat.xlsx")
kut.faturalari_ekstreyle_eslestir(EF, MT, mx)
wb = load_workbook(mx)
SONUC["Excel: mutabakat"] = [c.value for c in wb["Mutabakat"]["F"]][1:] == ["ödendi", "ödendi", "ödenmedi"]
ex = kut.ekstreyi_excele_aktar(MT, os.path.join(G, "ekstre.xlsx"))
SONUC["Excel: ekstre yürüyen bakiye"] = load_workbook(ex)["Hareketler"]["F4"].value == 226460.0

# ---- TCMB (bülten biçimi)
TCMB = """<?xml version="1.0" encoding="UTF-8"?>
<Tarih_Date Tarih="16.10.2026" Date="10/16/2026" Bulten_No="2026/196">
<Currency CrossOrder="0" Kod="USD" CurrencyCode="USD"><Unit>1</Unit><Isim>ABD DOLARI</Isim>
<CurrencyName>US DOLLAR</CurrencyName><ForexBuying>41.8456</ForexBuying><ForexSelling>41.9210</ForexSelling>
<BanknoteBuying>41.8163</BanknoteBuying><BanknoteSelling>41.9839</BanknoteSelling><CrossRateUSD/><CrossRateOther/></Currency>
<Currency CrossOrder="9" Kod="JPY" CurrencyCode="JPY"><Unit>100</Unit><Isim>JAPON YENİ</Isim>
<ForexBuying>27.5110</ForexBuying><ForexSelling>27.6932</ForexSelling><BanknoteBuying>27.3260</BanknoteBuying>
<BanknoteSelling>27.8790</BanknoteSelling></Currency>
<Currency CrossOrder="20" Kod="XDR" CurrencyCode="XDR"><Unit>1</Unit><Isim>ÖZEL ÇEKME HAKKI (SDR)</Isim>
<ForexBuying>56.9184</ForexBuying><ForexSelling/><BanknoteBuying/><BanknoteSelling/></Currency>
</Tarih_Date>"""
tarih, kurlar = b.tcmb_xml_coz(TCMB.encode("utf-8"))
SONUC["TCMB: bülten çözümü"] = tarih == "16.10.2026" and kurlar["USD"]["doviz_satis"] == 41.921
SONUC["TCMB: 100 birimlik kur 1 birime çevrildi"] = abs(kurlar["JPY"]["doviz_alis"] - 0.27511) < 1e-9
SONUC["TCMB: boş kur None"] = kurlar["XDR"]["doviz_satis"] is None
kut._kur_onbellek["16.10.2026"] = (tarih, kurlar)
SONUC["Kur Al (önbellek, TL=1)"] = kut.kur_al("usd", "16.10.2026") == 41.921 and kut.kur_al("TL", "16.10.2026") == 1.0
if "--canli" in sys.argv:
    bt, k = b.tcmb_kurlari()
    SONUC[f"TCMB canlı: {bt} USD {k['USD']['doviz_satis']}"] = k["USD"]["doviz_satis"] > 1
    bt2, _ = b.tcmb_kurlari("2026-10-18")  # pazar: önceki iş günü
    SONUC[f"TCMB tatil → önceki iş günü ({bt2})"] = bt2 in ("16.10.2026", "17.10.2026")

print("\n=========== BAĞLAYICILAR ===========")
for k_, v in SONUC.items():
    print(("✅" if v else "❌"), k_)
print(f"{sum(SONUC.values())}/{len(SONUC)} geçti")
sys.exit(0 if all(SONUC.values()) else 1)

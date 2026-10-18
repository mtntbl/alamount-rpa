# -*- coding: utf-8 -*-
"""
Türkiye'ye özel bağlayıcılar — .robot görevlerinde:   Library    AlamountBaglayicilar

  * e-Fatura / e-Arşiv (UBL-TR XML ya da ZIP) okuma, Excel'e aktarma
  * TCMB döviz kurları (günlük ve geçmiş tarih; tatilde önceki iş günü)
  * Banka ekstresi (MT940) okuma, Excel'e aktarma, faturalarla eşleştirme (mutabakat)
  * VKN, TCKN ve IBAN doğrulama

Genel bağlayıcılar (e-posta, Exchange/Outlook, veritabanı, SAP GUI, REST, FTP, Microsoft 365,
Salesforce) pakette hazır gelen RPA.* kütüphaneleridir; bkz. KULLANIM.md "Bağlayıcılar".
"""

import datetime as dt
import glob
import os
import re
import urllib.error
import urllib.request
import xml.etree.ElementTree as ET
import zipfile
from decimal import Decimal, InvalidOperation

# ====================================================================== doğrulama


def vkn_gecerli_mi(vkn):
    """10 haneli Vergi Kimlik Numarası kontrol hanesi doğrulaması (GİB algoritması)."""
    v = re.sub(r"\s", "", str(vkn))
    if not re.fullmatch(r"\d{10}", v):
        return False
    toplam = 0
    for i in range(9):
        t = (int(v[i]) + 9 - i) % 10
        p = (t * 2 ** (9 - i)) % 9
        if t != 0 and p == 0:
            p = 9
        toplam += p
    return (10 - toplam % 10) % 10 == int(v[9])


def tckn_gecerli_mi(tc):
    t = re.sub(r"\s", "", str(tc))
    if not re.fullmatch(r"[1-9]\d{10}", t):
        return False
    d = [int(c) for c in t]
    return (sum(d[0:9:2]) * 7 - sum(d[1:8:2])) % 10 == d[9] and sum(d[:10]) % 10 == d[10]


def iban_gecerli_mi(iban):
    """IBAN mod-97 doğrulaması; TR IBAN'ı 26 karakter olmalı."""
    i = re.sub(r"\s", "", str(iban)).upper()
    if not re.fullmatch(r"[A-Z]{2}\d{2}[A-Z0-9]{10,30}", i) or (i.startswith("TR") and len(i) != 26):
        return False
    sayi = "".join(str(int(c, 36)) for c in i[4:] + i[:4])
    return int(sayi) % 97 == 1


def _ondalik(metin):
    """'1.234,56' / '1234.56' / '1234,56' → Decimal."""
    if metin is None:
        return None
    m = str(metin).strip().replace(" ", "")
    if not m:
        return None
    if "," in m and "." in m:
        m = m.replace(".", "").replace(",", ".") if m.rfind(",") > m.rfind(".") else m.replace(",", "")
    elif "," in m:
        m = m.replace(",", ".")
    try:
        return Decimal(m)
    except InvalidOperation:
        return None


def _sayi(d):
    return float(d) if d is not None else None


# ====================================================================== e-Fatura (UBL-TR)


def _ad(el):
    return el.tag.rsplit("}", 1)[-1]


def _bul(el, yol):
    """Ad alanından bağımsız yol: 'AccountingSupplierParty/Party/PartyName/Name'."""
    if el is None:
        return None
    for parca in yol.split("/"):
        el = next((c for c in el if _ad(c) == parca), None)
        if el is None:
            return None
    return el


def _hepsi(el, ad):
    return [c for c in el if _ad(c) == ad] if el is not None else []


def _metin(el, yol=None):
    e = _bul(el, yol) if yol else el
    return (e.text or "").strip() if e is not None and e.text else ""


def _taraf(parti):
    p = _bul(parti, "Party")
    if p is None:
        return {}
    kimlik = {}
    for pid in _hepsi(p, "PartyIdentification"):
        i = _bul(pid, "ID")
        if i is not None and i.text:
            kimlik[(i.get("schemeID") or "").upper()] = i.text.strip()
    ad = _metin(p, "PartyName/Name")
    if not ad:
        ad = " ".join(x for x in (_metin(p, "Person/FirstName"), _metin(p, "Person/FamilyName")) if x)
    vkn = kimlik.get("VKN") or kimlik.get("TCKN") or ""
    return {"ad": ad, "vkn_tckn": vkn, "vkn_tckn_gecerli": vkn_gecerli_mi(vkn) if len(vkn) == 10 else tckn_gecerli_mi(vkn),
            "vergi_dairesi": _metin(p, "PartyTaxScheme/TaxScheme/Name"), "il": _metin(p, "PostalAddress/CityName"),
            "ilce": _metin(p, "PostalAddress/CitySubdivisionName"), "eposta": _metin(p, "Contact/ElectronicMail")}


def _vergiler(el):
    sonuc = []
    for tt in _hepsi(el, "TaxTotal") + _hepsi(el, "WithholdingTaxTotal"):
        for st in _hepsi(tt, "TaxSubtotal"):
            sonuc.append({"tur": _metin(st, "TaxCategory/TaxScheme/Name"),
                          "kod": _metin(st, "TaxCategory/TaxScheme/TaxTypeCode"),
                          "oran": _sayi(_ondalik(_metin(st, "Percent"))),
                          "matrah": _sayi(_ondalik(_metin(st, "TaxableAmount"))),
                          "tutar": _sayi(_ondalik(_metin(st, "TaxAmount"))),
                          "tevkifat": _ad(tt) == "WithholdingTaxTotal"})
    return sonuc


def efatura_xml_coz(veri, kaynak=""):
    """UBL-TR fatura XML'ini (bytes) sözlüğe çevirir."""
    kok = ET.fromstring(veri)
    if _ad(kok) != "Invoice":
        raise ValueError(f"UBL-TR fatura değil ({_ad(kok)}): {kaynak}")
    toplam = _bul(kok, "LegalMonetaryTotal")
    satirlar = []
    for s in _hepsi(kok, "InvoiceLine"):
        miktar = _bul(s, "InvoicedQuantity")
        kdv = next((v for v in _vergiler(s) if v["kod"] == "0015"), None)
        satirlar.append({"sira": _metin(s, "ID"), "urun": _metin(s, "Item/Name"),
                         "miktar": _sayi(_ondalik(miktar.text if miktar is not None else None)),
                         "birim": miktar.get("unitCode", "") if miktar is not None else "",
                         "birim_fiyat": _sayi(_ondalik(_metin(s, "Price/PriceAmount"))),
                         "tutar": _sayi(_ondalik(_metin(s, "LineExtensionAmount"))),
                         "kdv_orani": kdv["oran"] if kdv else None, "kdv_tutari": kdv["tutar"] if kdv else None})
    vergiler = _vergiler(kok)
    f = {
        "kaynak": kaynak, "fatura_no": _metin(kok, "ID"), "ettn": _metin(kok, "UUID"),
        "tarih": _metin(kok, "IssueDate"), "saat": _metin(kok, "IssueTime"),
        "senaryo": _metin(kok, "ProfileID"), "tip": _metin(kok, "InvoiceTypeCode"),
        "para_birimi": _metin(kok, "DocumentCurrencyCode") or "TRY",
        "kur": _sayi(_ondalik(_metin(kok, "PricingExchangeRate/CalculationRate"))),
        "satici": _taraf(_bul(kok, "AccountingSupplierParty")),
        "alici": _taraf(_bul(kok, "AccountingCustomerParty")),
        "mal_hizmet_toplam": _sayi(_ondalik(_metin(toplam, "LineExtensionAmount"))),
        "vergi_haric_toplam": _sayi(_ondalik(_metin(toplam, "TaxExclusiveAmount"))),
        "vergi_dahil_toplam": _sayi(_ondalik(_metin(toplam, "TaxInclusiveAmount"))),
        "iskonto_toplam": _sayi(_ondalik(_metin(toplam, "AllowanceTotalAmount"))),
        "odenecek_tutar": _sayi(_ondalik(_metin(toplam, "PayableAmount"))),
        "kdv_toplam": round(sum(v["tutar"] or 0 for v in vergiler if v["kod"] == "0015" and not v["tevkifat"]), 2),
        "tevkifat_toplam": round(sum(v["tutar"] or 0 for v in vergiler if v["tevkifat"]), 2),
        "vergiler": vergiler, "satirlar": satirlar,
        "notlar": [n.text.strip() for n in _hepsi(kok, "Note") if n.text and n.text.strip()],
        "siparis_no": _metin(kok, "OrderReference/ID"),
        "irsaliye_no": _metin(kok, "DespatchDocumentReference/ID"),
    }
    uyarilar = []
    satir_toplam = round(sum(s["tutar"] or 0 for s in satirlar), 2)
    if f["mal_hizmet_toplam"] is not None and abs(satir_toplam - f["mal_hizmet_toplam"]) > 0.01:
        uyarilar.append(f"Satır toplamı ({satir_toplam}) ≠ mal/hizmet toplamı ({f['mal_hizmet_toplam']})")
    if f["satici"].get("vkn_tckn") and not f["satici"]["vkn_tckn_gecerli"]:
        uyarilar.append(f"Satıcı VKN/TCKN geçersiz: {f['satici']['vkn_tckn']}")
    if f["vergi_haric_toplam"] is not None and f["vergi_dahil_toplam"] is not None:
        hesap = round(f["vergi_haric_toplam"] + f["kdv_toplam"], 2)
        if abs(hesap - f["vergi_dahil_toplam"]) > 0.05:
            uyarilar.append(f"Vergi hariç + KDV ({hesap}) ≠ vergi dahil toplam ({f['vergi_dahil_toplam']})")
    f["uyarilar"] = uyarilar
    return f


def efatura_oku(yol):
    """XML ya da ZIP (içinde bir ya da birden çok XML) → fatura sözlükleri listesi."""
    if zipfile.is_zipfile(yol):
        sonuc = []
        with zipfile.ZipFile(yol) as z:
            for ad in z.namelist():
                if ad.lower().endswith(".xml") and z.getinfo(ad).file_size < 50 * 1024 * 1024:
                    sonuc.append(efatura_xml_coz(z.read(ad), f"{os.path.basename(yol)}:{ad}"))
        return sonuc
    with open(yol, "rb") as f:
        return [efatura_xml_coz(f.read(), os.path.basename(yol))]


def efatura_klasoru_oku(klasor):
    faturalar, hatalar = [], []
    for yol in sorted(glob.glob(os.path.join(klasor, "*"))):
        if not yol.lower().endswith((".xml", ".zip")):
            continue
        try:
            faturalar += efatura_oku(yol)
        except Exception as e:
            hatalar.append({"dosya": os.path.basename(yol), "hata": str(e)})
    return faturalar, hatalar


# ====================================================================== TCMB kurları

TCMB_URL = "https://www.tcmb.gov.tr/kurlar"


def tcmb_xml_coz(veri):
    kok = ET.fromstring(veri)
    tarih = kok.get("Tarih", "")
    kurlar = {}
    for c in kok.iter("Currency"):
        kod = c.get("Kod") or c.get("CurrencyCode")
        birim = int(_ondalik(_metin(c, "Unit")) or 1)

        def k(ad):
            d = _ondalik(_metin(c, ad))
            return float(d / birim) if d is not None else None

        kurlar[kod] = {"kod": kod, "ad": _metin(c, "Isim"), "birim": birim, "doviz_alis": k("ForexBuying"),
                       "doviz_satis": k("ForexSelling"), "efektif_alis": k("BanknoteBuying"),
                       "efektif_satis": k("BanknoteSelling")}
    return tarih, kurlar


def tcmb_kurlari(tarih=None, zaman_asimi=20, geri_gun=7):
    """TCMB gösterge kurları (1 birim için TL). Tarih verilmezse bugünkü bülten; hafta sonu/tatil ise
    önceki iş günü. Döner: (bülten tarihi 'GG.AA.YYYY', {kod: {...}})."""
    if tarih in (None, "", "bugun", "bugün"):
        adaylar = [f"{TCMB_URL}/today.xml"]
        gun = dt.date.today()
    else:
        gun = tarih if isinstance(tarih, dt.date) else _tarih_coz(tarih)
        adaylar = []
    for i in range(geri_gun + 1):
        g = gun - dt.timedelta(days=i)
        adaylar.append(f"{TCMB_URL}/{g:%Y%m}/{g:%d%m%Y}.xml")
    son_hata = None
    for url in adaylar:
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "AlamountRPA"})
            with urllib.request.urlopen(req, timeout=zaman_asimi) as r:
                return tcmb_xml_coz(r.read())
        except urllib.error.HTTPError as e:
            son_hata = e
            if e.code != 404:
                raise
    raise RuntimeError(f"TCMB kuru bulunamadı ({gun:%d.%m.%Y} ve önceki {geri_gun} gün): {son_hata}")


def _tarih_coz(t):
    t = str(t).strip()
    for bicim in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y%m%d"):
        try:
            return dt.datetime.strptime(t, bicim).date()
        except ValueError:
            pass
    raise ValueError(f"Tarih anlaşılamadı: {t} (GG.AA.YYYY ya da YYYY-AA-GG)")


# ====================================================================== MT940


def _mt940_tarih(yymmdd):
    return dt.date(2000 + int(yymmdd[:2]), int(yymmdd[2:4]), int(yymmdd[4:6]))


_BAKIYE = re.compile(r"^(?P<ba>[CD])(?P<tarih>\d{6})(?P<pb>[A-Z]{3})(?P<tutar>[\d,\.]+)")
_HAREKET = re.compile(r"^(?P<valor>\d{6})(?P<kayit>\d{4})?(?P<ba>R?[CD])(?P<fon>[A-Z])?(?P<tutar>\d[\d,]*)"
                      r"(?P<tip>[NSF][A-Z0-9]{3})(?P<ref>[^/\n]*)(?://(?P<bankaref>[^\n]*))?(?:\n(?P<ek>.*))?",
                      re.S)


def _bakiye(deger):
    m = _BAKIYE.match(deger.strip())
    if not m:
        return None
    t = _ondalik(m["tutar"])
    return {"tarih": _mt940_tarih(m["tarih"]).isoformat(), "para_birimi": m["pb"],
            "tutar": float(-t if m["ba"] == "D" else t)}


def mt940_coz(metin):
    """MT940 ekstresi → [{hesap, acilis, kapanis, hareketler, tutarli}] (dosyada birden çok ekstre olabilir)."""
    metin = metin.replace("\r\n", "\n").replace("\r", "\n")
    alanlar = []
    for satir in metin.split("\n"):
        m = re.match(r"^:(\d{2}[A-Z]?):(.*)$", satir)
        if m:
            alanlar.append([m.group(1), m.group(2)])
        elif alanlar and satir.strip() and not satir.startswith(("{", "-}")):
            alanlar[-1][1] += "\n" + satir
    ekstreler, e, son = [], None, None
    for etiket, deger in alanlar:
        if etiket == "20":
            e = {"referans": deger.strip(), "hesap": "", "acilis": None, "kapanis": None, "hareketler": []}
            ekstreler.append(e)
        elif e is None:
            continue
        elif etiket == "25":
            e["hesap"] = deger.strip()
        elif etiket == "28C":
            e["ekstre_no"] = deger.strip()
        elif etiket in ("60F", "60M"):
            e["acilis"] = _bakiye(deger)
        elif etiket in ("62F", "62M"):
            e["kapanis"] = _bakiye(deger)
        elif etiket == "61":
            m = _HAREKET.match(deger.strip())
            if not m:
                raise ValueError(f"MT940 :61: satırı çözülemedi: {deger[:60]}")
            valor = _mt940_tarih(m["valor"])
            kayit = valor
            if m["kayit"]:
                kayit = dt.date(valor.year, int(m["kayit"][:2]), int(m["kayit"][2:]))
                if (kayit - valor).days > 180:
                    kayit = kayit.replace(year=kayit.year - 1)
                elif (valor - kayit).days > 180:
                    kayit = kayit.replace(year=kayit.year + 1)
            tutar = _ondalik(m["tutar"])
            borc = m["ba"] in ("D", "RC")
            son = {"tarih": kayit.isoformat(), "valor": valor.isoformat(), "yon": "borç" if borc else "alacak",
                   "tutar": float(-tutar if borc else tutar), "islem_tipi": m["tip"], "referans": m["ref"].strip(),
                   "banka_referansi": (m["bankaref"] or "").strip(), "aciklama": (m["ek"] or "").strip()}
            e["hareketler"].append(son)
        elif etiket == "86" and son is not None and e["hareketler"] and e["hareketler"][-1] is son:
            ack = re.sub(r"\s*\n\s*", " ", deger).strip()
            son["aciklama"] = (son["aciklama"] + " " + ack).strip()
    for e in ekstreler:
        pb = (e["acilis"] or e["kapanis"] or {}).get("para_birimi", "TRY")
        for h in e["hareketler"]:
            h["para_birimi"] = pb
        if e["acilis"] and e["kapanis"]:
            hesap = round(e["acilis"]["tutar"] + sum(h["tutar"] for h in e["hareketler"]), 2)
            e["tutarli"] = abs(hesap - e["kapanis"]["tutar"]) < 0.005
        else:
            e["tutarli"] = None
    return ekstreler


def mt940_oku(yol):
    ham = open(yol, "rb").read()
    for kod in ("utf-8", "cp1254", "iso-8859-9"):
        try:
            return mt940_coz(ham.decode(kod))
        except UnicodeDecodeError:
            continue
    return mt940_coz(ham.decode("latin-1"))


def faturalari_eslestir(faturalar, hareketler, yon="borç"):
    """Faturaları banka hareketleriyle eşleştirir: önce açıklamadaki fatura numarası, sonra aynı tutar.
    ``yon``: borç (biz ödedik, gelen faturalar) | alacak (bize ödendi, giden faturalar)."""
    kullanildi = set()
    sonuc = []
    adaylar = [(i, h) for i, h in enumerate(hareketler) if h["yon"] == yon]

    def norm(s):
        return re.sub(r"[^A-Z0-9]", "", str(s).upper())

    for f in faturalar:
        tutar = f.get("odenecek_tutar") or 0
        no = norm(f.get("fatura_no", ""))
        bulunan, yontem = None, ""
        for i, h in adaylar:
            if i not in kullanildi and no and no in norm(h["aciklama"] + h["referans"]):
                bulunan, yontem = (i, h), "fatura no"
                break
        ayni_pb = lambda h: h.get("para_birimi", "TRY") == (f.get("para_birimi") or "TRY")  # noqa: E731
        if not bulunan:
            for i, h in adaylar:
                if i not in kullanildi and ayni_pb(h) and abs(abs(h["tutar"]) - tutar) < 0.005:
                    bulunan, yontem = (i, h), "tutar"
                    break
        if bulunan:
            kullanildi.add(bulunan[0])
            h = bulunan[1]
            if ayni_pb(h):
                fark = round(abs(h["tutar"]) - tutar, 2)
                durum = "ödendi" if abs(fark) < 0.005 else ("eksik ödendi" if fark < 0 else "fazla ödendi")
            else:  # döviz faturası TL ödendi: tutar kur farkı nedeniyle karşılaştırılamaz
                fark, durum = None, f"ödendi ({h.get('para_birimi')}; kur farkını kontrol edin)"
            sonuc.append({"fatura_no": f.get("fatura_no"), "karsi_taraf": (f.get("satici") if yon == "borç" else
                                                                           f.get("alici") or {}).get("ad", ""),
                          "fatura_tutari": tutar, "odeme_tarihi": h["tarih"], "odenen": abs(h["tutar"]),
                          "fark": fark, "durum": durum, "eslesme": yontem, "aciklama": h["aciklama"]})
        else:
            sonuc.append({"fatura_no": f.get("fatura_no"), "karsi_taraf": (f.get("satici") if yon == "borç" else
                                                                           f.get("alici") or {}).get("ad", ""),
                          "fatura_tutari": tutar, "odeme_tarihi": "", "odenen": 0, "fark": -tutar,
                          "durum": "ödenmedi", "eslesme": "", "aciklama": ""})
    eslesmeyen = [h for i, h in adaylar if i not in kullanildi]
    return sonuc, eslesmeyen


# ====================================================================== Excel


def _excel_yaz(yol, sayfalar):
    """sayfalar: [(ad, başlıklar, satırlar)] — markalı başlık, filtre, sütun genişliği, sayı biçimi."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    for ad, basliklar, satirlar in sayfalar:
        ws = wb.create_sheet(ad[:31])
        ws.append(basliklar)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="022351")
            c.alignment = Alignment(vertical="center")
        for s in satirlar:
            ws.append(["; ".join(map(str, v)) if isinstance(v, list) else v for v in s])
        for i, b in enumerate(basliklar, 1):
            uzun = max([len(str(b))] + [len(str(s[i - 1])) for s in satirlar if s[i - 1] is not None] or [8])
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(60, max(10, uzun + 2))
            for r in range(2, len(satirlar) + 2):
                if isinstance(ws.cell(r, i).value, float):
                    ws.cell(r, i).number_format = "#,##0.00"
        ws.freeze_panes = "A2"
        if satirlar:
            ws.auto_filter.ref = ws.dimensions
    os.makedirs(os.path.dirname(os.path.abspath(yol)), exist_ok=True)
    wb.save(yol)
    return yol


def faturalari_excele_yaz(faturalar, yol, hatalar=()):
    fb = ["Fatura No", "ETTN", "Tarih", "Senaryo", "Tip", "Satıcı", "Satıcı VKN/TCKN", "Alıcı", "Alıcı VKN/TCKN",
          "Para Birimi", "Kur", "Mal/Hizmet", "KDV", "Tevkifat", "Vergi Dahil", "Ödenecek", "Uyarılar", "Kaynak"]
    fs = [[f["fatura_no"], f["ettn"], f["tarih"], f["senaryo"], f["tip"], f["satici"].get("ad"),
           f["satici"].get("vkn_tckn"), f["alici"].get("ad"), f["alici"].get("vkn_tckn"), f["para_birimi"], f["kur"],
           f["mal_hizmet_toplam"], f["kdv_toplam"], f["tevkifat_toplam"], f["vergi_dahil_toplam"],
           f["odenecek_tutar"], f["uyarilar"], f["kaynak"]] for f in faturalar]
    sb = ["Fatura No", "Sıra", "Ürün/Hizmet", "Miktar", "Birim", "Birim Fiyat", "Tutar", "KDV %", "KDV Tutarı"]
    ss = [[f["fatura_no"], s["sira"], s["urun"], s["miktar"], s["birim"], s["birim_fiyat"], s["tutar"],
           s["kdv_orani"], s["kdv_tutari"]] for f in faturalar for s in f["satirlar"]]
    sayfalar = [("Faturalar", fb, fs), ("Satırlar", sb, ss)]
    if hatalar:
        sayfalar.append(("Okunamayanlar", ["Dosya", "Hata"], [[h["dosya"], h["hata"]] for h in hatalar]))
    return _excel_yaz(yol, sayfalar)


# ====================================================================== Robot Framework kütüphanesi


class AlamountBaglayicilar:
    """Türkiye'ye özel bağlayıcılar: e-Fatura (UBL-TR), TCMB kurları, MT940 banka ekstresi, VKN/TCKN/IBAN."""

    ROBOT_LIBRARY_SCOPE = "GLOBAL"

    def __init__(self):
        self._kur_onbellek = {}

    def _log(self, mesaj):
        try:
            from robot.api import logger

            logger.info(mesaj)
        except Exception:
            print(mesaj)

    def vergi_numarasi_gecerli_mi(self, vkn):
        return vkn_gecerli_mi(vkn)

    def tc_kimlik_numarasi_gecerli_mi(self, tckn):
        return tckn_gecerli_mi(tckn)

    def iban_gecerli_mi(self, iban):
        return iban_gecerli_mi(iban)

    def e_fatura_oku(self, dosya):
        """UBL-TR XML ya da ZIP dosyasındaki faturaları okur (liste döner). Her fatura: fatura_no, ettn, tarih,
        senaryo, tip, satici{ad, vkn_tckn, vergi_dairesi, il}, alici{...}, para_birimi, mal_hizmet_toplam,
        kdv_toplam, tevkifat_toplam, vergi_dahil_toplam, odenecek_tutar, satirlar[...], uyarilar[...]."""
        faturalar = efatura_oku(dosya)
        for f in faturalar:
            self._log(f"e-Fatura {f['fatura_no']} · {f['satici'].get('ad')} · {f['odenecek_tutar']} "
                      f"{f['para_birimi']}{' · UYARI: ' + '; '.join(f['uyarilar']) if f['uyarilar'] else ''}")
        return faturalar

    def e_fatura_klasorunu_oku(self, klasor):
        """Klasördeki tüm .xml/.zip faturaları okur; okunamayanlar loga yazılır, iş durmaz."""
        faturalar, hatalar = efatura_klasoru_oku(klasor)
        for h in hatalar:
            self._log(f"Okunamadı: {h['dosya']} — {h['hata']}")
        self._log(f"{len(faturalar)} fatura okundu, {len(hatalar)} dosya okunamadı")
        return faturalar

    def e_faturalari_excele_aktar(self, kaynak, excel):
        """Klasördeki (ya da tek dosyadaki) faturaları 'Faturalar' ve 'Satırlar' sayfalı Excel'e yazar."""
        if os.path.isdir(kaynak):
            faturalar, hatalar = efatura_klasoru_oku(kaynak)
        else:
            faturalar, hatalar = efatura_oku(kaynak), []
        yol = faturalari_excele_yaz(faturalar, excel, hatalar)
        self._log(f"{len(faturalar)} fatura → {yol}")
        return yol

    def tcmb_kurlarini_al(self, tarih=None):
        """TCMB gösterge kurları (1 birim için TL; JPY gibi 100 birimlikler 1 birime çevrilir).
        Hafta sonu/tatilde önceki iş gününün bülteni döner. Sözlük: {USD: {doviz_alis, doviz_satis, ...}}."""
        anahtar = str(tarih or "bugun")
        if anahtar not in self._kur_onbellek:
            bulten, kurlar = tcmb_kurlari(tarih)
            self._log(f"TCMB bülteni {bulten}: {len(kurlar)} kur")
            self._kur_onbellek[anahtar] = (bulten, kurlar)
        return self._kur_onbellek[anahtar][1]

    def kur_al(self, doviz, tarih=None, tur="doviz_satis"):
        """Tek kur: ``Kur Al  USD``, ``Kur Al  EUR  15.10.2026  doviz_alis``.
        tur: doviz_alis | doviz_satis | efektif_alis | efektif_satis"""
        kod = str(doviz).upper()
        if kod in ("TRY", "TL"):
            return 1.0
        kurlar = self.tcmb_kurlarini_al(tarih)
        if kod not in kurlar:
            raise ValueError(f"TCMB bülteninde {kod} yok")
        deger = kurlar[kod][tur]
        if deger is None:
            raise ValueError(f"{kod} için {tur} yayımlanmamış")
        return deger

    def mt940_ekstresini_oku(self, dosya):
        """MT940 banka ekstresini okur. Tek ekstreyse sözlük, birden çoksa liste döner. Sözlük: hesap,
        acilis{tarih, tutar}, kapanis{...}, hareketler[{tarih, valor, yon, tutar, aciklama, referans}], tutarli."""
        ekstreler = mt940_oku(dosya)
        for e in ekstreler:
            self._log(f"Ekstre {e['hesap']}: {len(e['hareketler'])} hareket, açılış {e['acilis'] and e['acilis']['tutar']}"
                      f", kapanış {e['kapanis'] and e['kapanis']['tutar']}, bakiye tutarlı: {e['tutarli']}")
        return ekstreler[0] if len(ekstreler) == 1 else ekstreler

    def ekstreyi_excele_aktar(self, dosya, excel):
        ekstreler = mt940_oku(dosya)
        b = ["Hesap", "Tarih", "Valör", "Yön", "Tutar", "Bakiye", "İşlem Tipi", "Referans", "Açıklama"]
        satirlar = []
        for e in ekstreler:
            bakiye = e["acilis"]["tutar"] if e["acilis"] else 0.0
            for h in e["hareketler"]:
                bakiye = round(bakiye + h["tutar"], 2)
                satirlar.append([e["hesap"], h["tarih"], h["valor"], h["yon"], h["tutar"], bakiye, h["islem_tipi"],
                                 h["referans"], h["aciklama"]])
        return _excel_yaz(excel, [("Hareketler", b, satirlar)])

    def faturalari_ekstreyle_eslestir(self, faturalar, ekstre, excel=None, yon="borç"):
        """Fatura–banka mutabakatı. ``faturalar``: E-Fatura Oku sonucu (liste) ya da klasör yolu;
        ``ekstre``: MT940 dosyası ya da okunmuş ekstre. ``yon``: borç (gelen faturalar, biz ödedik) | alacak.
        Döner: [{fatura_no, karsi_taraf, fatura_tutari, odenen, fark, durum, odeme_tarihi, eslesme}]."""
        if isinstance(faturalar, str):
            faturalar = efatura_klasoru_oku(faturalar)[0] if os.path.isdir(faturalar) else efatura_oku(faturalar)
        if isinstance(ekstre, str):
            ekstre = mt940_oku(ekstre)
        ekstreler = ekstre if isinstance(ekstre, list) else [ekstre]
        hareketler = [h for e in ekstreler for h in e["hareketler"]]
        sonuc, eslesmeyen = faturalari_eslestir(faturalar, hareketler, yon)
        ozet = {}
        for s in sonuc:
            ozet[s["durum"]] = ozet.get(s["durum"], 0) + 1
        self._log("Mutabakat: " + ", ".join(f"{k}: {v}" for k, v in ozet.items()) +
                  f"; faturayla eşleşmeyen {yon} hareketi: {len(eslesmeyen)}")
        if excel:
            b = ["Fatura No", "Karşı Taraf", "Fatura Tutarı", "Ödenen", "Fark", "Durum", "Ödeme Tarihi", "Eşleşme",
                 "Banka Açıklaması"]
            s = [[r["fatura_no"], r["karsi_taraf"], r["fatura_tutari"], r["odenen"], r["fark"], r["durum"],
                  r["odeme_tarihi"], r["eslesme"], r["aciklama"]] for r in sonuc]
            e = [[h["tarih"], h["tutar"], h["aciklama"]] for h in eslesmeyen]
            _excel_yaz(excel, [("Mutabakat", b, s), ("Eşleşmeyen Hareketler", ["Tarih", "Tutar", "Açıklama"], e)])
            self._log(f"Mutabakat Excel'i: {excel}")
        return sonuc


_ADLAR = {
    "vergi_numarasi_gecerli_mi": "Vergi Numarası Geçerli Mi", "tc_kimlik_numarasi_gecerli_mi": "TC Kimlik Numarası Geçerli Mi",
    "iban_gecerli_mi": "IBAN Geçerli Mi", "e_fatura_oku": "E-Fatura Oku", "e_fatura_klasorunu_oku": "E-Fatura Klasörünü Oku",
    "e_faturalari_excele_aktar": "E-Faturaları Excel'e Aktar", "tcmb_kurlarini_al": "TCMB Kurlarını Al",
    "kur_al": "Kur Al", "mt940_ekstresini_oku": "MT940 Ekstresini Oku", "ekstreyi_excele_aktar": "Ekstreyi Excel'e Aktar",
    "faturalari_ekstreyle_eslestir": "Faturaları Ekstreyle Eşleştir",
}
for _metot, _robot_adi in _ADLAR.items():
    getattr(AlamountBaglayicilar, _metot).robot_name = _robot_adi
